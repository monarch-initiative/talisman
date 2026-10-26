"""Tabular GO over-representation, Mann-Whitney rank test, and grounded Claude summarization.

Supports workbooks (.xlsx) or TSV/CSV files containing gene-level correlation statistics and GO
term annotations (such as the "genes incldg GO" sheet produced in transcriptomic correlation
studies), followed by grounded LLM synthesis via Claude Code with deterministic citation checks.
"""

import collections
import csv
import json
import re
import statistics
from pathlib import Path
from typing import Any, Dict, Iterable, List, Set, Tuple, Union

from openpyxl import load_workbook
from scipy.stats import hypergeom, mannwhitneyu

from talisman.claude_code import ClaudeCodeClient

SHEET = "genes incldg GO"
ROOTS = {"GO:0003674", "GO:0008150", "GO:0005575"}
MIN_SIZE = 3
FDR = 0.05
SIGNIFICANT = 0.05
SUGGESTIVE = 0.25
ONTOLOGIES = ("BP", "MF", "CC")
ONTOLOGY_ALIASES = {
    "BP": "BP",
    "P": "BP",
    "BIOLOGICAL_PROCESS": "BP",
    "BIOLOGICAL PROCESS": "BP",
    "MF": "MF",
    "F": "MF",
    "MOLECULAR_FUNCTION": "MF",
    "MOLECULAR FUNCTION": "MF",
    "CC": "CC",
    "C": "CC",
    "CELLULAR_COMPONENT": "CC",
    "CELLULAR COMPONENT": "CC",
}
REQUIRED_COLUMNS = (
    "gene_id",
    "go_id",
    "go_ontology",
    "go_term",
    "Significant?",
    "pos/neg corr",
    "r (Pearson)",
    "FPKM average (samples)",
    "hypothetical protein?",
)
GO_ID = re.compile(r"GO:\d{7}")
ORA_FILE = {"up": "enrichment_pos.tsv", "down": "enrichment_neg.tsv"}

COLUMNS = [
    "go_id",
    "ontology",
    "name",
    "study_count",
    "background_count",
    "fold_enrichment",
    "p",
    "fdr",
    "significant",
    "representative",
    "same_genes_as",
    "mean_r",
    "mean_fpkm",
    "genes",
]

RANK_COLUMNS = [
    "go_id",
    "ontology",
    "name",
    "n_genes",
    "median_r",
    "auc",
    "direction",
    "p",
    "fdr",
    "significant",
    "representative",
    "same_genes_as",
    "genes",
]

DEFAULT_CONTEXT = """\
Two tests were run on GO terms annotating at least 3 genes (ontology roots excluded), with BH
correction within each ontology:
- over-representation: one-sided hypergeometric test of the significant genes in this direction
  against all annotated genes (study_count of background_count genes, fold enrichment);
- rank test: two-sided Mann-Whitney U test of the Pearson r of a term's genes against all other
  annotated genes (AUC > 0.5 means the term's genes rise with temperature, < 0.5 that they fall).
When several significant terms have exactly the same genes, only the most specific is listed."""

DEFAULT_SYSTEM_PROMPT = """\
You are a biologist summarizing GO enrichment results for collaborators who asked whether an LLM
can make sense of long GO term lists without hallucinating. Use only the GO terms in the tables as
evidence and cite their GO IDs. Themes must rest on terms from the "Significant" table; terms from
the "Suggestive" table may appear only in hypotheses, labelled as tentative. You may use background
biological knowledge to interpret the results, but say when a statement is interpretation rather
than a result. Gene names are not provided: do not name genes or operons as if they were measured.
Note when listed terms are parents of one another rather than independent findings, and keep claims
proportional to effect sizes and FDRs."""

SCHEMA: Dict[str, Any] = {
    "type": "object",
    "properties": {
        "summary": {"type": "string", "description": "3-5 sentences on what the results show."},
        "themes": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "theme": {"type": "string"},
                    "go_ids": {"type": "array", "items": {"type": "string"}},
                    "evidence": {
                        "type": "string",
                        "description": "Which tests and statistics support it.",
                    },
                    "interpretation": {"type": "string"},
                },
                "required": ["theme", "go_ids", "evidence", "interpretation"],
            },
        },
        "hypotheses": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "hypothesis": {"type": "string"},
                    "go_ids": {"type": "array", "items": {"type": "string"}},
                    "why_tentative": {"type": "string"},
                },
                "required": ["hypothesis", "go_ids", "why_tentative"],
            },
        },
        "caveats": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["summary", "themes", "hypotheses", "caveats"],
}


def bh(pvalues: List[float]) -> List[float]:
    """Benjamini-Hochberg adjusted p-values, in input order."""
    m = len(pvalues)
    order = sorted(range(m), key=lambda i: pvalues[i])
    adjusted = [1.0] * m
    running = 1.0
    for rank in range(m, 0, -1):
        i = order[rank - 1]
        running = min(running, pvalues[i] * m / rank)
        adjusted[i] = running
    return adjusted


def _iter_annotation_rows(path: Path, sheet: str = SHEET) -> Iterable[Tuple[Any, ...]]:
    """Yield raw rows from either an Excel workbook sheet or a delimited text file."""
    suffix = path.suffix.lower()
    if suffix in (".tsv", ".txt", ".csv"):
        delimiter = "," if suffix == ".csv" else "\t"
        with open(path, newline="", encoding="utf-8") as fh:
            for row in csv.reader(fh, delimiter=delimiter):
                yield tuple(row)
    else:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet not in wb.sheetnames:
                raise ValueError(
                    f"Worksheet {sheet!r} not found in {path}; available: {wb.sheetnames}"
                )
            yield from wb[sheet].iter_rows(values_only=True)
        finally:
            wb.close()


def read_annotations(
    path: Union[Path, str], sheet: str = SHEET
) -> Tuple[Dict[str, Dict[str, Any]], Dict[str, Tuple[str, str]], Dict[str, Set[str]]]:
    """Genes (with their statistics), GO terms and the genes annotated to each term."""
    rows = iter(_iter_annotation_rows(Path(path), sheet=sheet))
    header = next(rows, None)
    if header is None:
        raise ValueError(f"Empty annotation table: {path}")
    col = {str(h).strip(): i for i, h in enumerate(header) if h is not None and str(h).strip()}
    missing = [c for c in REQUIRED_COLUMNS if c not in col]
    if missing:
        raise ValueError(f"Missing required columns in {path}: {', '.join(missing)}")
    max_idx = max(col[c] for c in REQUIRED_COLUMNS)
    genes: Dict[str, Dict[str, Any]] = {}
    terms: Dict[str, Tuple[str, str]] = {}
    term_genes: Dict[str, Set[str]] = collections.defaultdict(set)
    for r in rows:
        if not r or len(r) <= max_idx:
            continue
        raw_gene, raw_go_id = r[col["gene_id"]], r[col["go_id"]]
        if raw_gene is None or raw_go_id is None:
            continue
        gene, go_id = str(raw_gene).strip(), str(raw_go_id).strip()
        if not gene or not go_id:
            continue
        raw_ont = str(r[col["go_ontology"]]).strip()
        ontology = ONTOLOGY_ALIASES.get(raw_ont.upper())
        if ontology is None:
            raise ValueError(
                f"Unknown GO ontology {raw_ont!r} for {go_id}; expected one of {ONTOLOGIES}"
            )
        genes[gene] = {
            "significant": str(r[col["Significant?"]]).strip().upper() == "YES",
            "sign": str(r[col["pos/neg corr"]]).strip(),
            "r": float(r[col["r (Pearson)"]]),
            "fpkm": float(r[col["FPKM average (samples)"]]),
            "hypothetical": str(r[col["hypothetical protein?"]]).strip().lower() == "yes",
        }
        terms[go_id] = (ontology, str(r[col["go_term"]]).strip())
        term_genes[go_id].add(gene)
    return genes, terms, term_genes


def enrich(
    direction: str,
    genes: Dict[str, Dict[str, Any]],
    terms: Dict[str, Tuple[str, str]],
    term_genes: Dict[str, Set[str]],
    min_size: int = MIN_SIZE,
    fdr: float = FDR,
) -> Dict[str, Any]:
    """Test every eligible term for over-representation among one direction's significant genes."""
    if direction not in ("pos", "neg"):
        raise ValueError(f"direction must be 'pos' or 'neg', got {direction!r}")
    study = {g for g, v in genes.items() if v["significant"] and v["sign"] == direction}
    big_n, n = len(genes), len(study)
    rows: List[Dict[str, Any]] = []
    for go_id, members in term_genes.items():
        if go_id in ROOTS or len(members) < min_size:
            continue
        hits = sorted(members & study)
        k, big_k = len(hits), len(members)
        ontology, name = terms[go_id]
        rows.append(
            {
                "go_id": go_id,
                "ontology": ontology,
                "name": name,
                "study_count": k,
                "background_count": big_k,
                "fold_enrichment": (k / n) / (big_k / big_n) if n and big_n else 0.0,
                "p": float(hypergeom.sf(k - 1, big_n, big_k, n)) if k else 1.0,
                "mean_r": statistics.mean(genes[g]["r"] for g in hits) if hits else None,
                "mean_fpkm": statistics.mean(genes[g]["fpkm"] for g in hits) if hits else None,
                "genes": hits,
                "representative": False,
                "same_genes_as": "",
            }
        )
    for ontology in ONTOLOGIES:
        sub = [r for r in rows if r["ontology"] == ontology]
        for r, q in zip(sub, bh([r["p"] for r in sub])):
            r["fdr"] = q
    for r in rows:
        r["significant"] = r["fdr"] <= fdr
    groups = collections.defaultdict(list)
    for r in rows:
        if r["significant"]:
            groups[(r["ontology"], tuple(r["genes"]))].append(r)
    for group in groups.values():
        group.sort(key=lambda r: (r["background_count"], r["p"], r["go_id"]))
        group[0]["representative"] = True
        for r in group[1:]:
            r["same_genes_as"] = group[0]["go_id"]
    rows.sort(key=lambda r: (r["fdr"], r["p"], r["go_id"]))
    return {"direction": direction, "study_size": n, "background_size": big_n, "rows": rows}


def rank_test(
    genes: Dict[str, Dict[str, Any]],
    terms: Dict[str, Tuple[str, str]],
    term_genes: Dict[str, Set[str]],
    min_size: int = MIN_SIZE,
    fdr: float = FDR,
) -> List[Dict[str, Any]]:
    """Mann-Whitney test of each eligible term's Pearson r against the other annotated genes."""
    rows: List[Dict[str, Any]] = []
    for go_id, members in term_genes.items():
        if go_id in ROOTS or len(members) < min_size or len(members) == len(genes):
            continue
        inside = [genes[g]["r"] for g in members]
        outside = [v["r"] for g, v in genes.items() if g not in members]
        u, p = mannwhitneyu(inside, outside, alternative="two-sided")
        auc = u / (len(inside) * len(outside))
        ontology, name = terms[go_id]
        rows.append(
            {
                "go_id": go_id,
                "ontology": ontology,
                "name": name,
                "n_genes": len(members),
                "median_r": statistics.median(inside),
                "auc": auc,
                "direction": "up" if auc > 0.5 else "down",
                "p": float(p),
                "genes": sorted(members),
                "representative": False,
                "same_genes_as": "",
            }
        )
    for ontology in ONTOLOGIES:
        sub = [r for r in rows if r["ontology"] == ontology]
        for r, q in zip(sub, bh([r["p"] for r in sub])):
            r["fdr"] = q
    for r in rows:
        r["significant"] = r["fdr"] <= fdr
    groups = collections.defaultdict(list)
    for r in rows:
        if r["significant"]:
            groups[(r["ontology"], tuple(r["genes"]))].append(r)
    for group in groups.values():
        group.sort(key=lambda r: (r["p"], r["go_id"]))
        group[0]["representative"] = True
        for r in group[1:]:
            r["same_genes_as"] = group[0]["go_id"]
    rows.sort(key=lambda r: (r["fdr"], r["p"], r["go_id"]))
    return rows


def write_tsv(rows: List[Dict[str, Any]], columns: List[str], path: Union[Path, str]) -> None:
    """Write enrichment or rank-test rows to a tab-separated file."""
    with open(path, "w", newline="", encoding="utf-8") as fh:
        out = csv.writer(fh, delimiter="\t")
        out.writerow(columns)
        for r in rows:
            values = {**r, "genes": ";".join(r["genes"])}
            for key in ("fold_enrichment", "mean_r", "mean_fpkm", "median_r", "auc"):
                if key in r:
                    values[key] = "" if r[key] is None else f"{r[key]:.4g}"
            for key in ("p", "fdr"):
                values[key] = f"{r[key]:.3g}"
            out.writerow([values[c] for c in columns])


def run_tabular_enrichment(
    workbook: Union[Path, str],
    outdir: Union[Path, str],
    sheet: str = SHEET,
    min_size: int = MIN_SIZE,
    fdr: float = FDR,
) -> Dict[str, Any]:
    """Run hypergeometric over-representation and Mann-Whitney rank tests and write TSV/JSON."""
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    genes, terms, term_genes = read_annotations(workbook, sheet=sheet)
    summary: Dict[str, Any] = {
        "workbook_sheet": sheet,
        "background_genes": len(genes),
        "annotated_hypothetical_genes": sum(v["hypothetical"] for v in genes.values()),
        "go_terms": len(terms),
        "settings": {
            "test": "one-sided hypergeometric",
            "correction": "BH within each ontology",
            "fdr": fdr,
            "min_background_genes": min_size,
            "roots_excluded": sorted(ROOTS),
        },
        "directions": {},
    }
    for direction in ("pos", "neg"):
        result = enrich(direction, genes, terms, term_genes, min_size=min_size, fdr=fdr)
        write_tsv(result["rows"], COLUMNS, outdir / f"enrichment_{direction}.tsv")
        rows = result["rows"]
        summary["directions"][direction] = {
            "study_genes": result["study_size"],
            "tested_terms": {o: sum(r["ontology"] == o for r in rows) for o in ONTOLOGIES},
            "significant_terms": {
                o: sum(r["significant"] and r["ontology"] == o for r in rows) for o in ONTOLOGIES
            },
            "representative_terms": {
                o: sum(r["representative"] and r["ontology"] == o for r in rows) for o in ONTOLOGIES
            },
        }
    ranked = rank_test(genes, terms, term_genes, min_size=min_size, fdr=fdr)
    write_tsv(ranked, RANK_COLUMNS, outdir / "rank_test.tsv")
    summary["rank_test"] = {
        "test": "two-sided Mann-Whitney U on Pearson r, term genes vs other annotated genes",
        "tested_terms": {o: sum(r["ontology"] == o for r in ranked) for o in ONTOLOGIES},
        "significant_terms": {
            d: {
                o: sum(
                    r["significant"] and r["direction"] == d and r["ontology"] == o for r in ranked
                )
                for o in ONTOLOGIES
            }
            for d in ("up", "down")
        },
        "representative_terms": {
            d: {
                o: sum(
                    r["representative"] and r["direction"] == d and r["ontology"] == o
                    for r in ranked
                )
                for o in ONTOLOGIES
            }
            for d in ("up", "down")
        },
    }
    (outdir / "enrichment_summary.json").write_text(
        json.dumps(summary, indent=2) + "\n", encoding="utf-8"
    )
    return summary


def read_tsv(path: Union[Path, str]) -> List[Dict[str, str]]:
    """Read a tab-separated file into a list of row dicts."""
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh, delimiter="\t"))


def collapse(rows: List[Dict[str, str]], size_key: str) -> List[Dict[str, str]]:
    """One row per distinct gene set: the most specific term, then the smallest p."""
    best: Dict[Tuple[str, str], Dict[str, str]] = {}
    for r in sorted(rows, key=lambda r: (int(r[size_key]), float(r["p"]), r["go_id"])):
        best.setdefault((r["ontology"], r["genes"]), r)
    return list(best.values())


def term_tables(
    results: Union[Path, str],
    direction: str,
    significant_cutoff: float = SIGNIFICANT,
    suggestive_cutoff: float = SUGGESTIVE,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Significant and suggestive terms for one direction, merged across both tests."""
    if direction not in ORA_FILE:
        raise ValueError(f"direction must be 'up' or 'down', got {direction!r}")
    results = Path(results)
    ora = read_tsv(results / ORA_FILE[direction])
    rank = [r for r in read_tsv(results / "rank_test.tsv") if r["direction"] == direction]
    terms: Dict[str, Dict[str, Any]] = {}
    for test, rows, size_key in (("ora", ora, "background_count"), ("rank", rank, "n_genes")):
        for tier, low, high in (
            ("significant", -1.0, significant_cutoff),
            ("suggestive", significant_cutoff, suggestive_cutoff),
        ):
            chosen = collapse([r for r in rows if low < float(r["fdr"]) <= high], size_key)
            for r in chosen:
                entry = terms.setdefault(
                    r["go_id"],
                    {
                        "go_id": r["go_id"],
                        "ontology": r["ontology"],
                        "name": r["name"],
                        "tier": tier,
                    },
                )
                if tier == "significant":
                    entry["tier"] = "significant"
                entry[test] = r
    significant = [t for t in terms.values() if t["tier"] == "significant"]
    suggestive = [t for t in terms.values() if t["tier"] == "suggestive"]

    def best_fdr(t: Dict[str, Any]) -> float:
        return min(float(t[k]["fdr"]) for k in ("ora", "rank") if k in t)

    return sorted(significant, key=best_fdr), sorted(suggestive, key=best_fdr)


def render(rows: List[Dict[str, Any]]) -> str:
    """Format merged term entries as a Markdown table."""
    lines = [
        "| GO ID | ontology | term | over-representation | rank test |",
        "|---|---|---|---|---|",
    ]
    for t in rows:
        ora = t.get("ora")
        rank = t.get("rank")
        ora_txt = (
            f"{ora['study_count']}/{ora['background_count']} genes, fold {ora['fold_enrichment']}, "
            f"FDR {ora['fdr']}"
            if ora
            else "-"
        )
        rank_txt = (
            f"{rank['n_genes']} genes, median r {rank['median_r']}, AUC {rank['auc']}, "
            f"FDR {rank['fdr']}"
            if rank
            else "-"
        )
        lines.append(f"| {t['go_id']} | {t['ontology']} | {t['name']} | {ora_txt} | {rank_txt} |")
    return "\n".join(lines)


def build_prompt(
    direction: str,
    significant: List[Dict[str, Any]],
    suggestive: List[Dict[str, Any]],
    context: str = DEFAULT_CONTEXT,
    significant_cutoff: float = SIGNIFICANT,
    suggestive_cutoff: float = SUGGESTIVE,
    condition: str = "temperature",
) -> str:
    """Build the user prompt for one direction's enrichment tables."""
    if direction not in ORA_FILE:
        raise ValueError(f"direction must be 'up' or 'down', got {direction!r}")
    label = "rise" if direction == "up" else "fall"
    sig_hdr = (
        f"Significant (FDR <= {significant_cutoff} in at least one test): "
        f"{len(significant)} terms"
    )
    sug_hdr = (
        f"Suggestive (FDR {significant_cutoff}-{suggestive_cutoff}, not significant): "
        f"{len(suggestive)} terms"
    )
    return (
        f"{context}\n\nSummarize the GO terms whose genes {label} with {condition}.\n\n"
        f"{sig_hdr}\n\n"
        f"{render(significant) if significant else '(none)'}\n\n"
        f"{sug_hdr}\n\n"
        f"{render(suggestive) if suggestive else '(none)'}\n"
    )


def ground(
    reply: Dict[str, Any],
    significant: List[Dict[str, Any]],
    suggestive: List[Dict[str, Any]],
) -> Dict[str, List[str]]:
    """Audit GO IDs cited in the structured reply against the significant and suggestive tables."""
    sig_ids = {t["go_id"] for t in significant}
    all_ids = sig_ids | {t["go_id"] for t in suggestive}
    explicit_theme_ids = {
        m
        for th in reply["themes"]
        for raw in th["go_ids"]
        for m in (GO_ID.findall(str(raw)) or [str(raw).strip()])
        if m
    }
    prose_theme_ids = set(GO_ID.findall(json.dumps(reply["themes"])))
    theme_ids = explicit_theme_ids | prose_theme_ids
    explicit_hyp_ids = {
        m
        for h in reply["hypotheses"]
        for raw in h["go_ids"]
        for m in (GO_ID.findall(str(raw)) or [str(raw).strip()])
        if m
    }
    cited = set(GO_ID.findall(json.dumps(reply))) | explicit_theme_ids | explicit_hyp_ids
    return {
        "cited_ids": sorted(cited),
        "ungrounded_ids": sorted(cited - all_ids),
        "theme_ids_not_significant": sorted(theme_ids - sig_ids),
    }


def to_markdown(direction: str, out: Dict[str, Any], condition: str = "temperature") -> str:
    """Format a structured tabular enrichment summary as a Markdown report."""
    if direction not in ORA_FILE:
        raise ValueError(f"direction must be 'up' or 'down', got {direction!r}")
    r = out["reply"]
    label = "rise" if direction == "up" else "fall"
    lines = [
        f"# GO terms whose genes {label} with {condition}",
        "",
        r["summary"],
        "",
        "## Themes",
        "",
    ]
    for th in r["themes"]:
        lines += [
            f"**{th['theme']}** ({', '.join(th['go_ids'])})",
            "",
            f"- Evidence: {th['evidence']}",
            f"- Interpretation: {th['interpretation']}",
            "",
        ]
    lines += ["## Hypotheses (tentative)", ""]
    for h in r["hypotheses"]:
        ids_str = ", ".join(h["go_ids"])
        lines += [f"- **{h['hypothesis']}** ({ids_str}). Why tentative: {h['why_tentative']}"]
    lines += ["", "## Caveats", ""] + [f"- {c}" for c in r["caveats"]]
    g = out["grounding"]
    cost = "unknown" if out["cost_usd"] is None else f"${out['cost_usd']:.3f}"
    sig_n = out["significant_terms"]
    sug_n = out["suggestive_terms"]
    lines += [
        "",
        "## Run",
        "",
        f"- Model {out['model']}, cost {cost}, cached {out['cached']}.",
        f"- Terms shown: {sig_n} significant, {sug_n} suggestive.",
        f"- Ungrounded IDs: {g['ungrounded_ids'] or 'none'}; theme IDs not significant: "
        f"{g['theme_ids_not_significant'] or 'none'}.",
    ]
    return "\n".join(lines) + "\n"


def summarize_tabular_enrichment(
    results: Union[Path, str],
    direction: str,
    client: ClaudeCodeClient,
    context: str = DEFAULT_CONTEXT,
    system_prompt: str = DEFAULT_SYSTEM_PROMPT,
    significant_cutoff: float = SIGNIFICANT,
    suggestive_cutoff: float = SUGGESTIVE,
    condition: str = "temperature",
) -> Dict[str, Any]:
    """Summarize one direction's tabular GO results with Claude and write JSON + Markdown."""
    if direction not in ORA_FILE:
        raise ValueError(f"direction must be 'up' or 'down', got {direction!r}")
    results = Path(results)
    significant, suggestive = term_tables(
        results,
        direction,
        significant_cutoff=significant_cutoff,
        suggestive_cutoff=suggestive_cutoff,
    )
    prompt = build_prompt(
        direction,
        significant,
        suggestive,
        context=context,
        significant_cutoff=significant_cutoff,
        suggestive_cutoff=suggestive_cutoff,
        condition=condition,
    )
    reply = client.complete(prompt, system_prompt, SCHEMA)
    out: Dict[str, Any] = {
        "direction": direction,
        "model": reply["model"],
        "cost_usd": reply["cost_usd"],
        "cached": reply["cached"],
        "significant_terms": len(significant),
        "suggestive_terms": len(suggestive),
        "grounding": ground(reply["structured_output"], significant, suggestive),
        "reply": reply["structured_output"],
        "prompt": prompt,
    }
    (results / f"summary_{direction}.json").write_text(
        json.dumps(out, indent=2) + "\n", encoding="utf-8"
    )
    (results / f"summary_{direction}.md").write_text(
        to_markdown(direction, out, condition=condition), encoding="utf-8"
    )
    return out
